"""EIP DDR Intelligence v4.4 — Export Routes: Excel + Professional PDF"""
import io, datetime, json, os
from fastapi import APIRouter, Request, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from core.database import load_latest, get_settings
from core.auth import require_permission, get_current_user
from core.audit import audit_from_request

router = APIRouter()

def _get_user_export(request: Request, token: str = Query(default=None)):
    """Accept token from query param (window.open) or Authorization header."""
    from core.auth import decode_token
    if token:
        try: return decode_token(token)
        except Exception: pass
    auth = request.headers.get("Authorization","")
    if auth.startswith("Bearer "):
        try: return decode_token(auth[7:])
        except Exception: pass
    raise HTTPException(401, "Authentication required for export.")

def _data(request):
    s = getattr(request.app.state,"intelligence_store",{})
    if s and s.get("opportunities"): return s
    return load_latest() or {}

@router.get("/excel")
async def export_excel(request: Request, token: str = Query(default=None), user: dict = Depends(_get_user_export)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    data = _data(request); settings = get_settings()
    co = settings.get("company_name","EnergiPro")
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    hf = Font(bold=True, color="FFFFFF", size=11)
    thin = Border(left=Side(style="thin"),right=Side(style="thin"),
                  top=Side(style="thin"),bottom=Side(style="thin"))
    urgency_colors = {"immediate":"FFCCCC","5_days":"FFE5B4","7_days":"CCE5FF",
                      "30_days":"E8F4FD","future":"F0F0F0"}

    # Sheet 1: Market Radar
    ws = wb.active; ws.title = "Market Radar"
    cols = ["Rank","Title","Product Line","Rig","Well","Field","Confidence %",
            "Est. Value ($)","Urgency","Competitor","Contact Role","Contact Name","Evidence","Action","Status"]
    for j,col in enumerate(cols,1):
        c=ws.cell(1,j,col); c.font=hf; c.fill=header_fill
        c.alignment=Alignment(horizontal="center")
    for i,opp in enumerate(data.get("opportunities",[]),2):
        row=[opp.get("rank",i-1),opp.get("title",""),opp.get("product_line",""),
             opp.get("rig",""),opp.get("well",""),opp.get("field",""),
             opp.get("confidence",0),round(opp.get("estimated_value",0)/1000)*1000,
             opp.get("urgency",""),opp.get("competitor",""),
             opp.get("suggested_contact",""),opp.get("contact_name",""),
             opp.get("evidence_text","")[:200],opp.get("action",""),opp.get("status","open")]
        fc=urgency_colors.get(opp.get("urgency","future"),"FFFFFF")
        for j,val in enumerate(row,1):
            c=ws.cell(i,j,val); c.border=thin
            if j<=10: c.fill=PatternFill("solid",fgColor=fc)
    col_widths=[6,42,20,12,16,14,12,14,12,16,28,20,50,44,10]
    for j,w in enumerate(col_widths,1): ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions

    # Sheet 2: Revenue
    ws2=wb.create_sheet("Revenue Forecast")
    ws2.append(["Product Line","Opportunities","Est. Value ($)","Top Rig","Competitor"])
    from collections import Counter
    pl_data={}
    for o in data.get("opportunities",[]):
        pl=o.get("product_line","Other")
        if pl not in pl_data: pl_data[pl]={"n":0,"v":0,"rigs":Counter(),"comps":Counter()}
        pl_data[pl]["n"]+=1; pl_data[pl]["v"]+=o.get("estimated_value",0)
        if o.get("rig"): pl_data[pl]["rigs"][o["rig"]]+=1
        if o.get("competitor"): pl_data[pl]["comps"][o["competitor"]]+=1
    for pl,d in sorted(pl_data.items(),key=lambda x:-x[1]["v"]):
        tr=d["rigs"].most_common(1)[0][0] if d["rigs"] else "—"
        tc=d["comps"].most_common(1)[0][0] if d["comps"] else "—"
        ws2.append([pl,d["n"],round(d["v"]),tr,tc])
    for cell in ws2[1]: cell.font=hf; cell.fill=header_fill
    ws2.column_dimensions["A"].width=24; ws2.column_dimensions["C"].width=16

    # Sheet 3: Competitors
    ws3=wb.create_sheet("Competitor Intelligence")
    ws3.append(["Company","Mentions","Rigs Active","Services"])
    cd={}
    for c in data.get("competitors",[]):
        name=c.get("normalized","?")
        if name not in cd: cd[name]={"n":0,"rigs":set(),"svcs":set()}
        cd[name]["n"]+=1
        if c.get("rig"): cd[name]["rigs"].add(c["rig"])
        if c.get("service_category"): cd[name]["svcs"].add(c["service_category"])
    for name,d in sorted(cd.items(),key=lambda x:-x[1]["n"]):
        ws3.append([name,d["n"],len(d["rigs"]),"; ".join(list(d["svcs"])[:3])])
    for cell in ws3[1]: cell.font=hf; cell.fill=header_fill
    ws3.column_dimensions["A"].width=22; ws3.column_dimensions["D"].width=40

    # Sheet 4: Actions
    from core.database import load_actions
    ws4=wb.create_sheet("Action Tracker")
    ws4.append(["Title","Owner","Due Date","Priority","Status","Expected Outcome","Notes"])
    for a in load_actions():
        ws4.append([a.get("title",""),a.get("owner",""),a.get("due_date",""),
                    a.get("priority",""),a.get("status",""),
                    a.get("expected_outcome",""),a.get("notes","")])
    for cell in ws4[1]: cell.font=hf; cell.fill=header_fill

    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    audit_from_request(request,"EXPORT_EXCEL",user=user,detail=f"Excel export: {co}")
    fname=f"{co.replace(' ','_')}_Intelligence_{datetime.date.today()}.xlsx"
    return StreamingResponse(buf,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition":f'attachment; filename="{fname}"'})


@router.get("/pdf-summary")
async def export_pdf_summary(request: Request, token: str = Query(default=None), user: dict = Depends(_get_user_export)):
    """Professional branded PDF briefing."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm, mm
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                         TableStyle, HRFlowable, KeepTogether)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    except ImportError:
        return {"error":"reportlab not installed"}

    data   = _data(request)
    settings = get_settings()
    co     = settings.get("company_name","EnergiPro")
    tagline= settings.get("company_tagline","Enterprise Oilfield Market Intelligence")
    today  = datetime.date.today().strftime("%d %B %Y")
    rdate  = data.get("report_date","—")
    kpis   = data.get("kpis",{})
    opps   = data.get("opportunities",[])
    comps  = data.get("competitors",[])
    actions= []
    try:
        from core.database import load_actions
        actions = load_actions()
    except: pass

    # ── Colours ──────────────────────────────────────────────────────────────
    NAVY   = colors.HexColor("#0D1F3C")
    TEAL   = colors.HexColor("#1F6460")
    BLUE   = colors.HexColor("#2563EB")
    AMBER  = colors.HexColor("#F59E0B")
    RED    = colors.HexColor("#EF4444")
    GREEN  = colors.HexColor("#10B981")
    LGREY  = colors.HexColor("#F4F6FB")
    MGREY  = colors.HexColor("#E2E8F0")
    WHITE  = colors.white
    BLACK  = colors.HexColor("#0F172A")

    def fv(v):
        if v>=1_000_000: return f"${v/1_000_000:.1f}M"
        if v>=1_000: return f"${v/1_000:.0f}k"
        return f"${v:.0f}"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=1.8*cm, rightMargin=1.8*cm,
                            topMargin=1.5*cm, bottomMargin=1.8*cm,
                            title=f"{co} – Intelligence Report")

    styles = getSampleStyleSheet()
    def sty(name,**kw):
        s=ParagraphStyle(name,parent=styles["Normal"],**kw)
        return s

    # ── Try to load real logo ─────────────────────────────────────────────────
    import os
    logo_path = os.path.join(os.path.dirname(__file__), '..', '..', 'logo-full.png')
    logo_path = os.path.normpath(logo_path)
    has_logo = os.path.exists(logo_path)

    S = {
        "h1":     sty("h1",fontName="Helvetica-Bold",fontSize=22,textColor=NAVY,spaceAfter=4,leading=26),
        "sub":    sty("sub",fontName="Helvetica",fontSize=10,textColor=colors.HexColor("#64748B"),spaceAfter=2),
        "h2":     sty("h2",fontName="Helvetica-Bold",fontSize=13,textColor=NAVY,spaceBefore=14,spaceAfter=6),
        "h3":     sty("h3",fontName="Helvetica-Bold",fontSize=10,textColor=TEAL,spaceBefore=6,spaceAfter=3),
        "body":   sty("body",fontName="Helvetica",fontSize=9,leading=14,textColor=BLACK,spaceAfter=4),
        "small":  sty("small",fontName="Helvetica",fontSize=8,textColor=colors.HexColor("#64748B")),
        "label":  sty("label",fontName="Helvetica-Bold",fontSize=7,textColor=colors.HexColor("#64748B")),
        "mono":   sty("mono",fontName="Courier",fontSize=8,leading=12,textColor=colors.HexColor("#334155")),
        "footer": sty("footer",fontName="Helvetica",fontSize=8,textColor=colors.HexColor("#94A3B8"),alignment=TA_CENTER),
        "right":  sty("right",fontName="Helvetica",fontSize=9,alignment=TA_RIGHT,textColor=colors.HexColor("#64748B")),
    }

    def hr(color=MGREY,thickness=0.5): return HRFlowable(width="100%",color=color,thickness=thickness,spaceAfter=8,spaceBefore=4)
    def sp(h=6): return Spacer(1,h)

    story = []

    # ── COVER HEADER ─────────────────────────────────────────────────────────
    if has_logo:
        from reportlab.platypus import Image as RLImage
        logo_img = RLImage(logo_path, width=3.2*cm, height=2.3*cm)
        header_data = [[
            logo_img,
            Paragraph(f"<b>Market Intelligence Report</b><br/>"
                      f"<font size=9>Report date: {rdate} &nbsp;·&nbsp; Generated: {today}</font>",
                      sty("hr2",fontName="Helvetica",fontSize=11,textColor=WHITE,alignment=TA_RIGHT)),
        ]]
    else:
        header_data = [[
            Paragraph(f"<b>{co}</b>", sty("hl",fontName="Helvetica-Bold",fontSize=20,textColor=WHITE)),
            Paragraph(f"<b>Market Intelligence Report</b><br/>"
                      f"<font size=9>Report date: {rdate} &nbsp;·&nbsp; Generated: {today}</font>",
                      sty("hr2",fontName="Helvetica",fontSize=11,textColor=WHITE,alignment=TA_RIGHT)),
        ]]
    ht = Table(header_data, colWidths=[9*cm, 8*cm])
    ht.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),NAVY),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(0,-1),16),
        ("RIGHTPADDING",(-1,0),(-1,-1),16),
        ("TOPPADDING",(0,0),(-1,-1),14),
        ("BOTTOMPADDING",(0,0),(-1,-1),14),
        ("ROUNDEDCORNERS",[4,4,4,4]),
    ]))
    story.append(ht); story.append(sp(12))

    # ── KPI SUMMARY BAR ──────────────────────────────────────────────────────
    imm_val = sum(o.get("estimated_value",0) for o in opps if o.get("urgency") in ("immediate","5_days"))
    kpi_items = [
        ("Open Opps", str(kpis.get("open_opportunities",0))),
        ("Pipeline", fv(kpis.get("pipeline_total",0))),
        ("7-Day Value", fv(imm_val)),
        ("Rigs Parsed", str(kpis.get("rigs_parsed",0))),
        ("Immediate", str(kpis.get("immediate_targets",0))),
        ("Competitors", str(kpis.get("competitor_mentions",0))),
    ]
    kpi_labels = [[Paragraph(k, sty("kl",fontName="Helvetica-Bold",fontSize=7,textColor=colors.HexColor("#64748B"),alignment=TA_CENTER)) for k,v in kpi_items]]
    kpi_values = [[Paragraph(v, sty("kv",fontName="Helvetica-Bold",fontSize=16,textColor=NAVY,alignment=TA_CENTER)) for k,v in kpi_items]]
    kt = Table(kpi_labels+kpi_values, colWidths=[2.85*cm]*6)
    kt.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),LGREY),
        ("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8),
        ("LINEAFTER",(0,0),(4,-1),0.5,MGREY),
        ("BOX",(0,0),(-1,-1),0.5,MGREY),
    ]))
    story.append(kt); story.append(sp(14))

    # ── EXECUTIVE NARRATIVE ──────────────────────────────────────────────────
    story.append(Paragraph("Executive Summary", S["h2"]))
    story.append(hr())
    narrative = data.get("narrative","") or ""
    for line in narrative.split("\n"):
        if line.strip():
            story.append(Paragraph(line.strip(), S["body"]))
    story.append(sp(8))

    # ── TOP OPPORTUNITIES TABLE ───────────────────────────────────────────────
    story.append(Paragraph("Top Commercial Opportunities", S["h2"]))
    story.append(hr())
    urg_map = {"immediate":"● NOW","5_days":"● 5d","7_days":"7d","30_days":"30d","future":"Future"}
    urg_col  = {"immediate":RED,"5_days":AMBER,"7_days":BLUE,"30_days":colors.HexColor("#6366F1"),"future":colors.HexColor("#94A3B8")}
    tbl_data = [[
        Paragraph("#",S["label"]),Paragraph("Product Line",S["label"]),
        Paragraph("Rig / Well",S["label"]),Paragraph("Conf",S["label"]),
        Paragraph("Value",S["label"]),Paragraph("Urgency",S["label"]),
        Paragraph("Competitor",S["label"]),Paragraph("Contact",S["label"]),
    ]]
    top_opps = opps[:20]
    for o in top_opps:
        urg = o.get("urgency","future")
        comp = o.get("competitor","") or "—"
        contact = (o.get("contact_name","") or o.get("suggested_contact","") or "—")[:22]
        tbl_data.append([
            Paragraph(str(o.get("rank","")), sty("c",fontName="Helvetica-Bold",fontSize=8,textColor=NAVY,alignment=TA_CENTER)),
            Paragraph(o.get("product_line",""), sty("c2",fontName="Helvetica",fontSize=8)),
            Paragraph(f"{o.get('rig','')} / {o.get('well','')}", sty("c3",fontName="Helvetica",fontSize=7,textColor=colors.HexColor("#334155"))),
            Paragraph(f"{o.get('confidence',0)}%", sty("c4",fontName="Helvetica-Bold",fontSize=8,alignment=TA_CENTER,textColor=GREEN if o.get('confidence',0)>=85 else BLUE)),
            Paragraph(fv(o.get("estimated_value",0)), sty("c5",fontName="Helvetica-Bold",fontSize=8,textColor=GREEN,alignment=TA_RIGHT)),
            Paragraph(urg_map.get(urg,"—"), sty("c6",fontName="Helvetica-Bold",fontSize=7,alignment=TA_CENTER,textColor=urg_col.get(urg,BLACK))),
            Paragraph(comp[:14], sty("c7",fontName="Helvetica",fontSize=7,textColor=RED if comp!="—" else colors.HexColor("#94A3B8"))),
            Paragraph(contact, sty("c8",fontName="Helvetica",fontSize=7,textColor=colors.HexColor("#334155"))),
        ])
    ot = Table(tbl_data, colWidths=[0.8*cm,3.2*cm,4.0*cm,1.3*cm,1.5*cm,1.5*cm,2.4*cm,2.3*cm])
    ot_style = [
        ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LGREY]),
        ("BOX",(0,0),(-1,-1),0.5,MGREY),("INNERGRID",(0,0),(-1,-1),0.25,MGREY),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
    ]
    # Highlight immediate rows
    for i,o in enumerate(top_opps,1):
        if o.get("urgency")=="immediate":
            ot_style.append(("LEFTPADDING",(0,i),(0,i),8))
            ot_style.append(("LINEAFTER",(0,i),(0,i),2,RED))
    ot.setStyle(TableStyle(ot_style))
    story.append(ot); story.append(sp(14))

    # ── COMPETITOR INTELLIGENCE ───────────────────────────────────────────────
    story.append(Paragraph("Competitor Presence", S["h2"]))
    story.append(hr())
    cd={}
    for c in comps:
        name=c.get("normalized","?")
        if name not in cd: cd[name]={"n":0,"rigs":set(),"svcs":set()}
        cd[name]["n"]+=1
        if c.get("rig"): cd[name]["rigs"].add(c["rig"])
        if c.get("service_category"): cd[name]["svcs"].add(c["service_category"])
    comp_rows=[[Paragraph("Company",S["label"]),Paragraph("Mentions",S["label"]),
                Paragraph("Rigs",S["label"]),Paragraph("Services",S["label"])]]
    for name,d in sorted(cd.items(),key=lambda x:-x[1]["n"])[:12]:
        comp_rows.append([
            Paragraph(f"<b>{name}</b>",sty("cn",fontName="Helvetica-Bold",fontSize=9,textColor=NAVY)),
            Paragraph(str(d["n"]),sty("cm",fontName="Helvetica-Bold",fontSize=9,textColor=RED,alignment=TA_CENTER)),
            Paragraph(str(len(d["rigs"])),sty("cr",fontName="Helvetica",fontSize=9,alignment=TA_CENTER)),
            Paragraph("; ".join(list(d["svcs"])[:2]),sty("cs",fontName="Helvetica",fontSize=8,textColor=colors.HexColor("#334155"))),
        ])
    ct=Table(comp_rows,colWidths=[4*cm,2.5*cm,2.5*cm,8*cm])
    ct.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),TEAL),("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LGREY]),
        ("BOX",(0,0),(-1,-1),0.5,MGREY),("INNERGRID",(0,0),(-1,-1),0.25,MGREY),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),
        ("LEFTPADDING",(0,0),(-1,-1),6),
    ]))
    story.append(ct); story.append(sp(14))

    # ── PRODUCT LINE DEMAND ───────────────────────────────────────────────────
    story.append(Paragraph("Pipeline by Product Line", S["h2"]))
    story.append(hr())
    pl_data={}
    for o in opps:
        pl=o.get("product_line","Other")
        if pl not in pl_data: pl_data[pl]={"n":0,"v":0}
        pl_data[pl]["n"]+=1; pl_data[pl]["v"]+=o.get("estimated_value",0)
    pl_rows=[[Paragraph("Product Line",S["label"]),Paragraph("Opps",S["label"]),
               Paragraph("Est. Value",S["label"]),Paragraph("% of Pipeline",S["label"])]]
    total_pipe=sum(d["v"] for d in pl_data.values()) or 1
    for pl,d in sorted(pl_data.items(),key=lambda x:-x[1]["v"]):
        pct=d["v"]/total_pipe*100
        pl_rows.append([
            Paragraph(pl,sty("pl",fontName="Helvetica",fontSize=9)),
            Paragraph(str(d["n"]),sty("pln",fontName="Helvetica-Bold",fontSize=9,alignment=TA_CENTER,textColor=BLUE)),
            Paragraph(fv(d["v"]),sty("plv",fontName="Helvetica-Bold",fontSize=9,textColor=GREEN,alignment=TA_RIGHT)),
            Paragraph(f"{pct:.0f}%",sty("plp",fontName="Helvetica",fontSize=9,textColor=colors.HexColor("#64748B"),alignment=TA_CENTER)),
        ])
    plt=Table(pl_rows,colWidths=[7*cm,2.5*cm,3*cm,4.5*cm])
    plt.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LGREY]),
        ("BOX",(0,0),(-1,-1),0.5,MGREY),("INNERGRID",(0,0),(-1,-1),0.25,MGREY),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),
        ("LEFTPADDING",(0,0),(-1,-1),6),
    ]))
    story.append(plt); story.append(sp(14))

    # ── ACTION TRACKER SUMMARY ────────────────────────────────────────────────
    if actions:
        story.append(Paragraph("Action Tracker Summary", S["h2"]))
        story.append(hr())
        by_s={"open":0,"contacted":0,"won":0,"lost":0}
        for a in actions: by_s[a.get("status","open")] = by_s.get(a.get("status","open"),0)+1
        at_data=[[
            Paragraph(f"<b>{len(actions)}</b><br/><font size=7>Total Actions</font>",sty("at",fontName="Helvetica-Bold",fontSize=14,textColor=NAVY,alignment=TA_CENTER)),
            Paragraph(f"<b>{by_s.get('contacted',0)}</b><br/><font size=7>Contacted</font>",sty("at2",fontName="Helvetica-Bold",fontSize=14,textColor=BLUE,alignment=TA_CENTER)),
            Paragraph(f"<b>{by_s.get('won',0)}</b><br/><font size=7>Won</font>",sty("at3",fontName="Helvetica-Bold",fontSize=14,textColor=GREEN,alignment=TA_CENTER)),
            Paragraph(f"<b>{by_s.get('lost',0)}</b><br/><font size=7>Lost</font>",sty("at4",fontName="Helvetica-Bold",fontSize=14,textColor=RED,alignment=TA_CENTER)),
            Paragraph(f"<b>{round(by_s.get('won',0)/max(len(actions),1)*100)}%</b><br/><font size=7>Win Rate</font>",sty("at5",fontName="Helvetica-Bold",fontSize=14,textColor=TEAL,alignment=TA_CENTER)),
        ]]
        att=Table(at_data,colWidths=[3.4*cm]*5)
        att.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),LGREY),
            ("TOPPADDING",(0,0),(-1,-1),10),("BOTTOMPADDING",(0,0),(-1,-1),10),
            ("LINEAFTER",(0,0),(3,-1),0.5,MGREY),("BOX",(0,0),(-1,-1),0.5,MGREY)]))
        story.append(att); story.append(sp(8))
        # List open actions
        open_actions=[a for a in actions if a.get("status","open")=="open"][:8]
        if open_actions:
            story.append(Paragraph("Open Actions", S["h3"]))
            for a in open_actions:
                story.append(Paragraph(f"<b>{a.get('title','')}</b> · {a.get('owner','')} · Due: {a.get('due_date','—')} · <i>{a.get('expected_outcome','')[:80]}</i>", S["small"]))
        story.append(sp(14))

    # ── FOOTER ───────────────────────────────────────────────────────────────
    story.append(hr(MGREY,1))
    story.append(Paragraph(
        f"<b>{co}</b> – {tagline} &nbsp;|&nbsp; "
        f"EIP DDR Intelligence v4.4 &nbsp;|&nbsp; "
        f"Confidential — For authorised recipients only &nbsp;|&nbsp; Generated {today}",
        S["footer"]))
    story.append(Paragraph(
        "Generated by EIP DDR Market Intelligence &amp; Opportunity Radar — "
        "Values reflect Aramco unit-rate service contracts per job, not full well costs. "
        "This report is intended for internal commercial use only.",
        sty("disc",fontName="Helvetica",fontSize=7,
            textColor=colors.HexColor("#94A3B8"),alignment=TA_CENTER,spaceBefore=2)))

    doc.build(story)
    buf.seek(0)
    audit_from_request(request,"EXPORT_PDF",user=user,detail=f"PDF brief: {co}")
    fname=f"{co.replace(' ','_')}_Intelligence_Brief_{datetime.date.today()}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition":f'attachment; filename="{fname}"'})


@router.get("/pdf-evidence")
async def export_pdf_evidence(
    request: Request,
    token: str = Query(default=None),
    validation_status: str = None,
    product_line: str = None,
    min_confidence: int = 65,
    user: dict = Depends(_get_user_export)
):
    """
    PDF Opportunity Evidence Report — detailed per-opportunity evidence,
    confidence breakdown, validation status, and source traceability.
    Suitable for handing to a technical review team.
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                         TableStyle, HRFlowable, KeepTogether, PageBreak)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    except ImportError:
        return {"error": "reportlab not installed"}

    data     = _data(request)
    settings = get_settings()
    co       = settings.get("company_name", "EnergiPro")
    today    = __import__("datetime").date.today().strftime("%d %B %Y")
    opps     = data.get("opportunities", [])

    # Apply filters
    if validation_status: opps = [o for o in opps if o.get("validation_status") == validation_status]
    if product_line:      opps = [o for o in opps if o.get("product_line","") == product_line]
    if min_confidence:    opps = [o for o in opps if o.get("confidence",0) >= min_confidence]

    NAVY  = colors.HexColor("#0D1F3C")
    TEAL  = colors.HexColor("#1F6460")
    BLUE  = colors.HexColor("#2563EB")
    RED   = colors.HexColor("#EF4444")
    GREEN = colors.HexColor("#10B981")
    AMBER = colors.HexColor("#F59E0B")
    LGREY = colors.HexColor("#F4F6FB")
    MGREY = colors.HexColor("#E2E8F0")
    WHITE = colors.white

    def fv(v):
        if v >= 1_000_000: return f"${v/1_000_000:.1f}M"
        if v >= 1_000:     return f"${v/1_000:.0f}k"
        return f"${v:.0f}"

    def sty(name, **kw):
        s = ParagraphStyle(name, parent=getSampleStyleSheet()["Normal"], **kw)
        return s

    def hr(color=MGREY, thickness=0.5):
        return HRFlowable(width="100%", color=color, thickness=thickness,
                           spaceAfter=6, spaceBefore=4)

    buf = __import__("io").BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=1.8*cm, rightMargin=1.8*cm,
                            topMargin=1.5*cm, bottomMargin=1.8*cm,
                            title=f"{co} — Opportunity Evidence Report")
    story = []
    import json as _json

    # Cover
    filters_applied = []
    if validation_status: filters_applied.append(f"Validation: {validation_status}")
    if product_line:      filters_applied.append(f"Product Line: {product_line}")
    if min_confidence > 0:filters_applied.append(f"Min Confidence: {min_confidence}%")
    filter_str = " | ".join(filters_applied) if filters_applied else "All opportunities"

    story.append(Paragraph(f"<b>{co}</b>", sty("h1",fontName="Helvetica-Bold",fontSize=22,textColor=NAVY)))
    story.append(Paragraph("Opportunity Evidence Report", sty("h2",fontName="Helvetica",fontSize=14,textColor=TEAL)))
    story.append(Paragraph(f"Generated: {today} | Filters: {filter_str} | {len(opps)} opportunities",
                            sty("sub",fontName="Helvetica",fontSize=10,textColor=colors.HexColor("#64748B"),spaceAfter=12)))
    story.append(hr(NAVY, 1))
    story.append(Spacer(1, 12))

    urg_map   = {"immediate":"● IMMEDIATE","5_days":"● 5 DAYS","7_days":"7 DAYS","30_days":"30 DAYS","future":"FUTURE"}
    urg_color = {"immediate":RED,"5_days":AMBER,"7_days":BLUE,"30_days":colors.HexColor("#6366F1"),"future":colors.HexColor("#94A3B8")}
    val_color = {"accepted":GREEN,"rejected":RED,"needs_review":AMBER,"converted":colors.HexColor("#8B5CF6"),"new":colors.HexColor("#64748B")}

    from engines.ingestion_pipeline import compute_confidence_breakdown

    for i, opp in enumerate(opps[:50], 1):  # Max 50 in evidence report
        try:
            kws = _json.loads(opp.get("matched_keywords","[]") or "[]")
        except Exception:
            kws = []
        opp["matched_keywords"] = kws

        urg = opp.get("urgency","future")
        val = opp.get("validation_status","new")
        conf = opp.get("confidence",0)
        comp = opp.get("competitor","")
        breakdown = compute_confidence_breakdown(opp)

        block = []

        # Header row
        hdr_data = [[
            Paragraph(f"<b>#{i} &nbsp; {opp.get('title','')}</b>",
                      sty("oh",fontName="Helvetica-Bold",fontSize=11,textColor=NAVY)),
            Paragraph(f"<b>{urg_map.get(urg,'—')}</b>",
                      sty("ou",fontName="Helvetica-Bold",fontSize=10,textColor=urg_color.get(urg,NAVY),alignment=TA_RIGHT)),
        ]]
        ht = Table(hdr_data, colWidths=[13*cm, 4*cm])
        ht.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,-1),LGREY),
            ("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8),
            ("LEFTPADDING",(0,0),(0,-1),10),
            ("LINEBELOW",(0,0),(-1,-1),1,NAVY),
        ]))
        block.append(ht)
        block.append(Spacer(1,6))

        # Meta grid
        meta_data = [[
            Paragraph(f"<b>Rig:</b> {opp.get('rig','—')}",         sty("m",fontName="Helvetica",fontSize=9)),
            Paragraph(f"<b>Well:</b> {opp.get('well','—')}",        sty("m",fontName="Helvetica",fontSize=9)),
            Paragraph(f"<b>Field:</b> {opp.get('field','—')}",      sty("m",fontName="Helvetica",fontSize=9)),
            Paragraph(f"<b>Competitor:</b> {comp or 'None detected'}",
                      sty("mc",fontName="Helvetica",fontSize=9,textColor=RED if comp else colors.HexColor("#64748B"))),
        ],[
            Paragraph(f"<b>Product Line:</b> {opp.get('product_line','—')}", sty("m",fontName="Helvetica",fontSize=9)),
            Paragraph(f"<b>Est. Value:</b> {fv(opp.get('estimated_value',0))}",
                      sty("mv",fontName="Helvetica-Bold",fontSize=9,textColor=GREEN)),
            Paragraph(f"<b>Confidence:</b> {conf}% — {breakdown.get('confidence_category','')}",
                      sty("mconf",fontName="Helvetica",fontSize=9,
                          textColor=GREEN if conf>=80 else AMBER if conf>=65 else RED)),
            Paragraph(f"<b>Status:</b> {val.replace('_',' ').title()}",
                      sty("ms",fontName="Helvetica-Bold",fontSize=9,textColor=val_color.get(val,NAVY))),
        ]]
        mt = Table(meta_data, colWidths=[4.25*cm]*4)
        mt.setStyle(TableStyle([
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
            ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("LINEAFTER",(0,0),(2,-1),0.3,MGREY),
        ]))
        block.append(mt)
        block.append(Spacer(1,6))

        # Evidence box
        ev_text = (opp.get("evidence_text","") or "")[:400]
        src_info = ""
        if opp.get("source_file"):
            src_info = f" | Source: {opp.get('source_file','')} p.{opp.get('source_page',0) or '?'}"
        block.append(Paragraph(
            f"<b>Evidence ({opp.get('evidence_section','DDR')}){src_info}:</b>",
            sty("el",fontName="Helvetica-Bold",fontSize=8,textColor=colors.HexColor("#64748B"),spaceAfter=2)))
        block.append(Paragraph(ev_text or "No evidence captured.",
            sty("ev",fontName="Courier",fontSize=8,leading=12,
                textColor=colors.HexColor("#334155"),
                leftIndent=8,borderPad=4)))
        block.append(Spacer(1,4))

        # Confidence breakdown
        block.append(Paragraph("<b>Confidence Breakdown</b>",
            sty("cb",fontName="Helvetica-Bold",fontSize=8,textColor=colors.HexColor("#64748B"),spaceAfter=2)))
        factors = [
            ("Product Line Match", breakdown.get("pl_match",0), 25),
            ("Activity Signal",    breakdown.get("activity",0), 20),
            ("Rig/Well Clarity",   breakdown.get("rig_clarity",0), 15),
            ("Competitor Signal",  breakdown.get("competitor",0), 15),
            ("Evidence Quality",   breakdown.get("evidence",0), 15),
            ("Recency",            breakdown.get("recency",0), 10),
        ]
        row1 = [Paragraph(f"<b>{n}</b><br/>{v}/{mx}", sty(f"cf{j}",fontName="Helvetica",fontSize=7,alignment=TA_CENTER)) for j,(n,v,mx) in enumerate(factors[:3])]
        row2 = [Paragraph(f"<b>{n}</b><br/>{v}/{mx}", sty(f"cf{j}",fontName="Helvetica",fontSize=7,alignment=TA_CENTER)) for j,(n,v,mx) in enumerate(factors[3:])]
        ct = Table([row1,row2], colWidths=[5.67*cm]*3)
        ct.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,-1),LGREY),
            ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("LINEAFTER",(0,0),(1,-1),0.3,MGREY),
            ("LINEBEFORE",(1,0),(2,-1),0.3,MGREY),
            ("LINEBELOW",(0,0),(-1,0),0.3,MGREY),
        ]))
        block.append(ct)
        block.append(Spacer(1,4))

        # Recommended action
        block.append(Paragraph(
            f"<b>Action:</b> {opp.get('action','—')} &nbsp;|&nbsp; <b>Contact:</b> {opp.get('suggested_contact','—')}",
            sty("act",fontName="Helvetica",fontSize=8,textColor=TEAL)))

        # Validation notes if any
        if opp.get("validator_comment") or opp.get("validated_by"):
            block.append(Paragraph(
                f"<b>Reviewer:</b> {opp.get('validated_by','—')} &nbsp;|&nbsp; {opp.get('validator_comment','')}",
                sty("rev",fontName="Helvetica",fontSize=8,textColor=colors.HexColor("#64748B"))))

        block.append(Spacer(1,10))
        block.append(hr())

        story.append(KeepTogether(block))

    # Footer
    story.append(Spacer(1,10))
    story.append(hr(MGREY,1))
    story.append(Paragraph(
        f"{co} — Opportunity Evidence Report | EIP DDR Intelligence v4.4 | {today} | "
        f"Generated by EIP DDR Market Intelligence &amp; Opportunity Radar | Confidential",
        sty("foot",fontName="Helvetica",fontSize=7,textColor=colors.HexColor("#94A3B8"),alignment=TA_CENTER)))

    doc.build(story)
    buf.seek(0)
    audit_from_request(request, "EXPORT_PDF", user=user,
                       detail=f"Evidence report: {len(opps)} opps | {filter_str}")
    fname = f"{co.replace(' ','_')}_Evidence_Report_{__import__('datetime').date.today()}.pdf"
    return __import__("fastapi.responses",fromlist=["StreamingResponse"]).StreamingResponse(
        buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ── PPTX Presentation Builder (v6.3) ─────────────────────────────────────────
@router.get("/presentation")
async def export_presentation(
    request: Request,
    token: str = Query(default=None),
    user: dict = Depends(_get_user_export)
):
    """
    Generate a 9-slide branded executive PowerPoint presentation.
    Uses live market intelligence data from current intelligence store.
    """
    from fastapi.responses import StreamingResponse as SR
    from engines.intelligence.presentation_builder import build_presentation
    from core.database import get_settings

    settings     = get_settings()
    co           = settings.get("company_name", "EnergiPro Solutions")
    store        = getattr(request.app.state, "intelligence_store", {})
    report_date  = store.get("report_date", "")

    buf   = build_presentation(store, company_name=co, report_date=report_date)
    fname = f"{co.replace(' ','_')}_Intelligence_{__import__('datetime').date.today()}.pptx"

    audit_from_request(request, "EXPORT_PDF", user=user,
                       detail="Executive presentation exported")
    return SR(buf,
              media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
              headers={"Content-Disposition": f'attachment; filename="{fname}"'})
    # NOTE: an exact duplicate of this route (same path+method, dead code —
    # FastAPI only ever matches the first registration) was removed here
    # during the enterprise-demo port.
